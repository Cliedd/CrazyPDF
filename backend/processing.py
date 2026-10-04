import io
import json
import math
import subprocess
import tempfile
import os
from pathlib import Path
import fitz
from PIL import Image, ImageEnhance, ImageOps, ImageDraw
from backend.presets import PRESETS, dimensions

Image.MAX_IMAGE_PIXELS = 40_000_000
_rembg_session = None
os.environ.setdefault('U2NET_HOME', str(Path('data/models').resolve()))
os.environ.setdefault('TESSDATA_PREFIX', str(Path('data/tessdata').resolve()))
os.environ.setdefault('OMP_NUM_THREADS', '1')

def remove_background(image):
    global _rembg_session
    from rembg import new_session, remove
    if _rembg_session is None:
        _rembg_session = new_session('u2netp')
    return remove(image.convert('RGBA'), session=_rembg_session)

def load_image(path):
    image = Image.open(path)
    if image.width * image.height > Image.MAX_IMAGE_PIXELS:
        raise ValueError('Image trop grande : maximum 40 millions de pixels.')
    image = ImageOps.exif_transpose(image)
    if image.info.get('icc_profile'):
        from PIL import ImageCms
        try:
            image = ImageCms.profileToProfile(image, ImageCms.ImageCmsProfile(io.BytesIO(image.info['icc_profile'])), ImageCms.createProfile('sRGB'), outputMode='RGB')
        except Exception:
            raise ValueError('Profil de couleur non reconnu. Importez une image sRGB.')
    return image.convert('RGBA')

def encode_image(image, target, fmt='JPEG', dpi=300, max_kb=None):
    if fmt == 'JPEG':
        if image.mode != 'RGB':
            canvas = Image.new('RGB', image.size, 'white')
            canvas.paste(image, mask=image.getchannel('A') if image.mode == 'RGBA' else None)
            image = canvas
        for quality in range(95, 9, -5):
            buffer = io.BytesIO()
            image.save(buffer, format='JPEG', quality=quality, optimize=True, dpi=(dpi, dpi))
            if not max_kb or len(buffer.getvalue()) <= max_kb * 1000:
                target.write_bytes(buffer.getvalue())
                return
        raise ValueError('Cette photo ne peut pas respecter la taille maximale sans trop dégrader la qualité. Essayez un cadrage plus simple.')
    image.save(target, format=fmt, dpi=(dpi, dpi))

def photo(source, output, options):
    preset = PRESETS[options['preset']]
    if preset.get('no_retouch') and (options.get('remove_bg') or options.get('brightness', 0) != 0):
        raise ValueError('Cette démarche interdit les retouches numériques : utilisez le recadrage sans détourage ni exposition.')
    image = load_image(source)
    if options.get('remove_bg'):
        image = remove_background(image)
    width, height = dimensions(preset)
    background = options.get('background') or preset['background']
    scale = max(width / image.width, height / image.height) * float(options.get('zoom', 1))
    image = image.resize((max(1, round(image.width * scale)), max(1, round(image.height * scale))), Image.Resampling.LANCZOS)
    image = image.rotate(-float(options.get('rotation', 0)), Image.Resampling.BICUBIC, expand=True)
    if options.get('brightness'):
        image = ImageEnhance.Brightness(image).enhance(1 + options['brightness'] / 100)
    canvas = Image.new('RGBA', (width, height), background)
    x = round((width - image.width) / 2 + float(options.get('offset_x', 0)) * width)
    y = round((height - image.height) / 2 + float(options.get('offset_y', 0)) * height)
    if preset.get('no_retouch'):
        coverage = Image.new('L', (width, height), 0)
        coverage.paste(image.getchannel('A'), (x, y))
        if coverage.getextrema()[0] < 255:
            raise ValueError('Le cadrage crée des zones artificielles : recentrez la photo ou augmentez le zoom. Cette démarche exige une photo originale sans modification du fond.')
    canvas.paste(image, (x, y), image)
    if options.get('export_type') == 'sheet':
        # Real 10×15 cm PDF page, actual-size images and crop marks.
        count = int(options.get('sheet_count', 4))
        page_width, page_height = 100 / 25.4 * 72, 150 / 25.4 * 72
        pw, ph = preset['width_mm'] / 25.4 * 72, preset['height_mm'] / 25.4 * 72
        cols = 2 if pw * 2 + 12 <= page_width else 1
        rows = math.ceil(count / cols)
        if ph * rows + 12 > page_height:
            raise ValueError('Ce format ne tient pas sur une planche 10 × 15 cm. Réduisez le nombre de photos.')
        doc = fitz.open()
        page = doc.new_page(width=page_width, height=page_height)
        stream = io.BytesIO()
        canvas.convert('RGB').save(stream, 'JPEG', quality=95)
        for i in range(count):
            left = (page_width - cols * pw) / 2 + i % cols * pw
            top = (page_height - rows * ph) / 2 + i // cols * ph
            rect = fitz.Rect(left, top, left + pw, top + ph)
            page.insert_image(rect, stream=stream.getvalue())
            for cx, cy in [(left, top), (left + pw, top), (left, top + ph), (left + pw, top + ph)]:
                page.draw_line((cx - 2, cy), (cx + 2, cy), color=(.5, .5, .5), width=.2)
                page.draw_line((cx, cy - 2), (cx, cy + 2), color=(.5, .5, .5), width=.2)
        doc.save(output)
        doc.close()
    else:
        fmt = options.get('format', 'JPEG')
        if preset.get('max_kb'):
            fmt = 'JPEG'
        encode_image(canvas, output, fmt, preset['dpi'], preset.get('max_kb'))
    return {'width_px': width, 'height_px': height, 'dpi': preset['dpi'], 'preset': preset['label'], 'export_type': options.get('export_type', 'single'), 'status': 'Format exporté — conformité visuelle à vérifier', 'source': preset['source']}

def cutout(source, output, options):
    image = remove_background(load_image(source))
    if options.get('sharpness'):
        image = ImageEnhance.Sharpness(image).enhance(1.3)
    if options.get('exposure'):
        alpha = image.getchannel('A')
        image = ImageOps.autocontrast(image.convert('RGB')).convert('RGBA')
        image.putalpha(alpha)
    if options.get('shadow'):
        image = ImageEnhance.Brightness(image).enhance(1.05)
    background = options.get('background', 'transparent')
    if background != 'transparent':
        canvas = Image.new('RGBA', image.size, background)
        canvas.alpha_composite(image)
        image = canvas
    encode_image(image, output, 'PNG')
    return {'width_px': image.width, 'height_px': image.height, 'status': 'Fond traité — vérifier les contours', 'background': background}

def office_to_pdf(source, output):
    with tempfile.TemporaryDirectory() as td:
        result = subprocess.run([os.getenv('LIBREOFFICE_BIN', 'libreoffice'), '-env:UserInstallation=file://' + td + '/profile', '--headless', '--convert-to', 'pdf', '--outdir', td, str(source.resolve())], capture_output=True, timeout=120)
        pdf = Path(td) / (source.stem + '.pdf')
        if result.returncode or not pdf.exists():
            import logging
            logging.getLogger('docuvisa').error('LibreOffice exit=%s stdout=%s stderr=%s', result.returncode, result.stdout.decode(errors='replace')[:1000], result.stderr.decode(errors='replace')[:1000])
            raise ValueError('LibreOffice ne peut pas convertir ce document. Vérifiez qu’il est valide et non protégé.')
        output.write_bytes(pdf.read_bytes())
    with fitz.open(output) as doc:
        return {'pages': len(doc), 'status': 'Converti avec succès'}

def convert(source, output, mode, options):
    if mode in ('word-pdf', 'pptx-pdf'):
        return office_to_pdf(source, output)
    if mode == 'pdf-word':
        from pdf2docx import Converter
        with fitz.open(source) as document:
            if document.needs_pass:
                raise ValueError('PDF protégé : fournissez une copie déverrouillée.')
            if len(document) > 200:
                raise ValueError('Maximum 200 pages par conversion.')
            scans = any(not page.get_text().strip() for page in document)
        actual = source
        with tempfile.TemporaryDirectory() as td:
            if scans:
                from docx import Document
                result = Document()
                with fitz.open(source) as document:
                    for index, page in enumerate(document):
                        if index:
                            result.add_page_break()
                        if page.get_text().strip():
                            text = page.get_text()
                        else:
                            tp = page.get_textpage_ocr(language='fra+eng', dpi=200, full=True, tessdata=os.environ['TESSDATA_PREFIX'])
                            text = page.get_text(textpage=tp)
                        for paragraph in text.split('\n'):
                            if paragraph.strip():
                                result.add_paragraph(paragraph.strip())
                result.save(output)
                return {'status': 'Texte OCR modifiable', 'ocr': True, 'note': 'Reconnaissance Tesseract français/anglais. Relisez le texte ; la mise en page des scans est reconstruite simplement.'}
            converter = Converter(str(actual))
            try:
                converter.convert(str(output), multi_processing=False)
            finally:
                converter.close()
        return {'status': 'Document Word modifiable', 'ocr': scans, 'note': 'Vérifiez les tableaux, formules et la mise en page après conversion.'}
    if mode == 'pdf-pptx':
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        with fitz.open(source) as document:
            if document.needs_pass:
                raise ValueError('PDF protégé : fournissez une copie déverrouillée.')
            if len(document) > 200 or len(document) == 0:
                raise ValueError('Le PDF doit contenir entre 1 et 200 pages.')
            first = document[0].rect
            presentation.slide_width = Inches(first.width / 72)
            presentation.slide_height = Inches(first.height / 72)
            for page in document:
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                slide.shapes.add_picture(io.BytesIO(pixmap.tobytes('png')), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
            presentation.save(output)
            return {'pages': len(document), 'status': 'Converti en diapositives', 'note': 'Chaque page devient une image haute résolution dans une diapositive ; les textes ne sont pas modifiables.'}
    raise ValueError('Conversion inconnue.')
