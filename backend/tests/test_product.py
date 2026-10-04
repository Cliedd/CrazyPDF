import io
import json
import os
import time
import uuid
import zipfile
from pathlib import Path
import tempfile

test_root = Path(tempfile.mkdtemp(prefix='docuvisa-tests-'))
os.environ['DATABASE_URL'] = 'sqlite:///' + str(test_root / 'test.db')
os.environ['STORAGE_DIR'] = str(test_root / 'files')
os.environ['INTERNAL_API_TOKEN'] = 'test-internal-token'
os.environ['WORKER_COUNT'] = '2'

import pytest
import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from pptx import Presentation
from fastapi.testclient import TestClient
from backend.main import app, Session, Job
from backend.presets import PRESETS, dimensions

@pytest.fixture(scope='module')
def client():
    with TestClient(app, headers={'X-Internal-Token': 'test-internal-token'}) as client:
        yield client

def account(client, suffix=''):
    email = str(uuid.uuid4()) + suffix + '@example.com'
    response = client.post('/api/auth/register', json={'email': email, 'password': 'Testing-Password-2026', 'name': 'Étudiante Test'})
    assert response.status_code == 200, response.text
    return email

def submit(client, filename, content, mode, options=None):
    response = client.post('/api/jobs', files={'file': (filename, content)}, data={'mode': mode, 'options': json.dumps(options or {})})
    assert response.status_code == 202, response.text
    job = response.json()
    deadline = time.time() + 180
    while time.time() < deadline:
        job = client.get('/api/jobs/' + job['id']).json()
        if job['status'] in ('completed', 'failed'):
            assert job['status'] == 'completed', job
            result = client.get(job['download_url'])
            assert result.status_code == 200
            return job, result.content
        time.sleep(.3)
    pytest.fail('Job did not complete')

def image_bytes(width=1800, height=2400):
    image = Image.new('RGB', (width, height), '#dedede')
    draw = ImageDraw.Draw(image)
    draw.ellipse((width*.3, height*.15, width*.7, height*.55), fill='#ae7760')
    draw.rectangle((width*.2, height*.6, width*.8, height), fill='#233552')
    buffer = io.BytesIO()
    image.save(buffer, 'JPEG', quality=90)
    return buffer.getvalue()

def test_auth_private_files_and_no_delete(client):
    assert client.get('/api/jobs').status_code == 401
    email = account(client)
    assert client.get('/api/auth/me').json()['email'] == email
    assert client.patch('/api/auth/me', json={'name': 'Nouveau nom'}).json()['name'] == 'Nouveau nom'
    _, content = submit(client, 'portrait.jpg', image_bytes(), 'photo', {'preset': 'campus-cm'})
    job = client.get('/api/jobs').json()[0]
    assert client.delete('/api/jobs/' + job['id']).status_code == 405
    assert client.post('/api/auth/logout').status_code == 200
    assert client.get(job['download_url']).status_code == 401
    assert client.post('/api/auth/login', json={'email': email, 'password': 'wrong-password-123'}).status_code == 401
    assert client.post('/api/auth/login', json={'email': email, 'password': 'Testing-Password-2026'}).status_code == 200
    first_cookies = dict(client.cookies)
    account(client, 'other')
    assert client.get('/api/jobs/' + job['id']).status_code == 404
    assert client.get(job['download_url']).status_code == 404
    client.cookies.clear()
    client.cookies.update(first_cookies)
    # Simulate Render replacing its ephemeral filesystem during a deployment.
    with Session() as session:
        stored = session.get(Job, job['id'])
        Path(stored.output).unlink()
    restored = client.get(job['download_url'])
    assert restored.status_code == 200
    assert restored.content == content

def test_official_preset_dimensions_and_limits(client):
    for preset_id, preset in PRESETS.items():
        job, content = submit(client, 'photo.jpg', image_bytes(), 'photo', {'preset': preset_id})
        image = Image.open(io.BytesIO(content))
        assert image.size == dimensions(preset)
        assert image.mode == 'RGB'
        if preset.get('max_kb'):
            assert len(content) <= preset['max_kb'] * 1000
        assert abs(image.info['dpi'][0] - preset['dpi']) <= 1
        assert job['metadata']['source'] == preset['source']
    for preset_id in ('us-visa', 'ca-pr', 'de-visa'):
        response = client.post('/api/jobs', files={'file': ('photo.jpg', image_bytes())}, data={'mode': 'photo', 'options': json.dumps({'preset': preset_id, 'remove_bg': True})})
        assert response.status_code == 422

def test_print_sheets(client):
    for count in (4, 6):
        job, content = submit(client, 'portrait.jpg', image_bytes(), 'photo', {'preset': 'fr-visa', 'export_type': 'sheet', 'sheet_count': count})
        with fitz.open(stream=content, filetype='pdf') as doc:
            assert abs(doc[0].rect.width / 72 * 25.4 - 100) < .01
            assert abs(doc[0].rect.height / 72 * 25.4 - 150) < .01
            assert len({image[0] for image in doc[0].get_images()}) == 1
            placements = doc[0].get_image_rects(doc[0].get_images()[0][0])
            assert len(placements) == count
            assert abs(placements[0].width / 72 * 25.4 - 35) < .01
            assert abs(placements[0].height / 72 * 25.4 - 45) < .01

def test_word_and_pdf_conversions(client):
    document = Document()
    document.add_heading('DocuVisa test', 0)
    document.add_paragraph('Lettre de motivation pour des études en France.')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Université'
    table.cell(0, 1).text = 'Formation'
    buffer = io.BytesIO()
    document.save(buffer)
    _, pdf = submit(client, 'lettre.docx', buffer.getvalue(), 'word-pdf')
    with fitz.open(stream=pdf, filetype='pdf') as result:
        assert 'DocuVisa test' in result[0].get_text()
    _, docx = submit(client, 'lettre.pdf', pdf, 'pdf-word')
    reopened = Document(io.BytesIO(docx))
    assert 'DocuVisa test' in ' '.join(p.text for p in reopened.paragraphs)
def test_pdf_to_pptx_conversion(client):
    doc = fitz.open()
    for text in ('Soutenance DocuVisa', 'Deuxième diapositive'):
        page = doc.new_page()
        page.insert_text((72, 72), text)
    source = doc.tobytes()
    doc.close()
    _, pptx = submit(client, 'soutenance.pdf', source, 'pdf-pptx')
    result = Presentation(io.BytesIO(pptx))
    assert len(result.slides) == 2
    assert len(result.slides[0].shapes) == 1

def test_pptx_to_pdf_conversion_when_impress_is_installed(client):
    if not Path('/usr/lib/libreoffice/program/libsdlo.so').exists():
        pytest.skip('LibreOffice Impress is installed in the Render Docker image, but not on this test host.')
    presentation = Presentation()
    for text in ('Soutenance DocuVisa', 'Deuxième diapositive'):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = text
    buffer = io.BytesIO()
    presentation.save(buffer)
    _, pdf = submit(client, 'soutenance.pptx', buffer.getvalue(), 'pptx-pdf')
    with fitz.open(stream=pdf, filetype='pdf') as result:
        assert len(result) == 2
        assert 'Soutenance DocuVisa' in result[0].get_text()

def test_scan_ocr_real_text(client):
    image = Image.new('RGB', (1400, 700), 'white')
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 60)
    ImageDraw.Draw(image).text((100, 200), 'Bonjour Campus France 2026', font=font, fill='black')
    buffer = io.BytesIO()
    image.save(buffer, 'PNG')
    doc = fitz.open()
    page = doc.new_page(width=700, height=350)
    page.insert_image(page.rect, stream=buffer.getvalue())
    content = doc.tobytes()
    doc.close()
    job, result = submit(client, 'scan.pdf', content, 'pdf-word')
    text = ' '.join(p.text for p in Document(io.BytesIO(result)).paragraphs)
    assert 'Campus France' in text, text
    assert job['metadata']['ocr'] is True

def test_real_cutout_and_archive_zip(client):
    source = Path('design/stitch/a9678693ab304b88b1be7d76bb882c71/screen.png').read_bytes()
    for background in ('transparent', '#FFFFFF', '#F1F5F9', '#E0F2FE'):
        _, content = submit(client, 'portrait.png', source, 'cutout', {'background': background, 'sharpness': True, 'exposure': True})
        image = Image.open(io.BytesIO(content))
        assert image.size == Image.open(io.BytesIO(source)).size
        if background == 'transparent':
            assert image.mode == 'RGBA'
            assert image.getchannel('A').getextrema()[0] < 10
            assert image.getchannel('A').getextrema()[1] > 200
        else:
            assert image.getpixel((0, 0))[:3] in [(255, 255, 255), (241, 245, 249), (224, 242, 254)]
    result = client.get('/api/jobs/export.zip')
    assert result.status_code == 200
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        assert len(archive.namelist()) > 10
        assert archive.testzip() is None

def test_invalid_files_and_options(client):
    for filename, content, mode, options in [('invalid.pdf', b'not a pdf', 'pdf-word', {}), ('empty.docx', b'', 'word-pdf', {}), ('invalid.docx', b'PKfake', 'word-pdf', {}), ('photo.jpg', image_bytes(), 'photo', {'preset': 'invented'}), ('photo.jpg', image_bytes(), 'photo', {'zoom': 100}), ('photo.jpg', image_bytes(), 'word-pdf', {})]:
        response = client.post('/api/jobs', files={'file': (filename, content)}, data={'mode': mode, 'options': json.dumps(options)})
        assert response.status_code == 422, response.text
    assert client.post('/api/auth/logout', headers={'Origin': 'https://untrusted.example'}).status_code == 403
    assert client.get('/api/jobs', headers={'X-Internal-Token': 'wrong'}).status_code == 403
