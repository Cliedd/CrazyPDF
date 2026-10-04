"""Focused conversion checks independent of the API job scheduler."""
import io
import os
import shutil
from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from pptx import Presentation
import pytest

from backend.processing import convert


@pytest.mark.skipif(not shutil.which(os.environ.get('LIBREOFFICE_BIN', 'libreoffice')), reason='LibreOffice unavailable')
def test_real_word_pdf_word_roundtrip(tmp_path):
    source, pdf, output = tmp_path / 'letter.docx', tmp_path / 'letter.pdf', tmp_path / 'letter-back.docx'
    document = Document()
    document.add_heading('DocuVisa university transcript', 0)
    document.add_paragraph('Lettre de motivation France Canada 2026')
    document.save(source)
    metadata = convert(source, pdf, 'word-pdf', {})
    assert metadata['pages'] == 1
    with fitz.open(pdf) as result:
        assert 'DocuVisa university transcript' in result[0].get_text()
    convert(pdf, output, 'pdf-word', {})
    assert 'DocuVisa university transcript' in ' '.join(p.text for p in Document(output).paragraphs)


@pytest.mark.skipif(not Path('/usr/lib/libreoffice/program/libsdlo.so').exists(), reason='LibreOffice Impress unavailable')
def test_real_presentation_pdf_presentation_roundtrip(tmp_path):
    source, pdf, output = tmp_path / 'slides.pptx', tmp_path / 'slides.pdf', tmp_path / 'slides-back.pptx'
    presentation = Presentation()
    for title in ('DocuVisa presentation', 'Deuxieme slide'):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
    presentation.save(source)
    metadata = convert(source, pdf, 'pptx-pdf', {})
    assert metadata['pages'] == 2
    with fitz.open(pdf) as result:
        assert 'DocuVisa presentation' in result[0].get_text()
        assert 'Deuxieme slide' in result[1].get_text()
    convert(pdf, output, 'pdf-pptx', {})
    assert len(Presentation(output).slides) == 2


def test_mixed_pdf_page_ratios_are_preserved(tmp_path):
    source = tmp_path / 'mixed.pdf'
    output = tmp_path / 'mixed.pptx'
    with fitz.open() as pdf:
        for width, height in ((600, 800), (800, 400), (7200, 3600)):
            page = pdf.new_page(width=width, height=height)
            page.insert_text((50, 50), 'Visible page content')
        pdf.save(source)
    metadata = convert(source, output, 'pdf-pptx', {})
    assert metadata['pages'] == 3
    presentation = Presentation(output)
    for slide, expected_ratio in zip(presentation.slides, (.75, 2, 2)):
        picture = slide.shapes[0]
        assert abs(picture.width / picture.height - expected_ratio) < .001
        assert picture.left >= 0 and picture.top >= 0
        assert picture.left + picture.width <= presentation.slide_width
        assert picture.top + picture.height <= presentation.slide_height
        with Image.open(io.BytesIO(picture.image.blob)) as image:
            assert image.width * image.height <= 4_010_000


def test_mixed_selectable_and_scanned_pdf_keeps_both_pages(tmp_path):
    source, output = tmp_path / 'mixed.pdf', tmp_path / 'mixed.docx'
    image = Image.new('RGB', (1400, 700), 'white')
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 60)
    ImageDraw.Draw(image).text((100, 200), 'Bonjour Campus France 2026', font=font, fill='black')
    stream = io.BytesIO()
    image.save(stream, 'PNG')
    with fitz.open() as pdf:
        page = pdf.new_page(width=700, height=350)
        page.insert_text((50, 80), 'Selectable university transcript')
        page = pdf.new_page(width=700, height=350)
        page.insert_image(page.rect, stream=stream.getvalue())
        pdf.save(source)
    metadata = convert(source, output, 'pdf-word', {})
    paragraphs = ' '.join(p.text for p in Document(output).paragraphs)
    assert metadata['ocr'] is True
    assert 'Selectable university transcript' in paragraphs
    assert 'Campus France 2026' in paragraphs
