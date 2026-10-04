from pathlib import Path
from pptx import Presentation
from backend.processing import office_to_pdf

root = Path('/tmp/docuvisa-office-test')
root.mkdir(exist_ok=True)
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text = 'DocuVisa PowerPoint'
presentation.save(root / 'presentation.pptx')
print(office_to_pdf(root / 'presentation.pptx', root / 'presentation.pdf'))
