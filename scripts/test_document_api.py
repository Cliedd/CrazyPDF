"""Real multipart/browser-session smoke test. Creates one technical account.

TEST_ORIGIN=https://... .venv/bin/python scripts/test_document_api.py
Install httpx[http2] to additionally exercise HTTP/2 on the live Render endpoint.
"""
import importlib.util
import io
import json
import os
import secrets
import time
import uuid

import fitz
import httpx
from docx import Document
from PIL import Image
from pptx import Presentation


ORIGIN = os.environ.get('TEST_ORIGIN', 'http://127.0.0.1:10001').rstrip('/')
HTTP2 = importlib.util.find_spec('h2') is not None
TIMEOUT = float(os.environ.get('TEST_JOB_TIMEOUT', '240'))


def expect(response, status):
    assert response.status_code == status, f'{response.request.method} {response.request.url.path}: HTTP {response.status_code}'
    return response


def assert_session(client, identity):
    user = expect(client.get('/api/auth/session'), 200).json()['user']
    assert user and user['id'] == identity, 'Session lost during conversion/navigation'
    # A new HTTP client represents a reload/new feature page with browser cookies.
    with httpx.Client(base_url=ORIGIN, cookies=client.cookies, timeout=60, http2=HTTP2) as reloaded:
        user = expect(reloaded.get('/api/auth/session'), 200).json()['user']
        assert user and user['id'] == identity, 'Session lost after page reload'


def submit(client, identity, filename, content, mode):
    assert_session(client, identity)
    start = time.monotonic()
    response = expect(client.post('/api/jobs', files={'file': (filename, content)}, data={'mode': mode, 'options': '{}'}), 202)
    job = response.json()
    assert_session(client, identity)
    deadline = time.monotonic() + TIMEOUT
    while time.monotonic() < deadline:
        job = expect(client.get('/api/jobs/' + job['id']), 200).json()
        assert_session(client, identity)
        if job['status'] == 'failed':
            raise AssertionError(f'{mode}: processing failed: {job.get("error")}')
        if job['status'] == 'completed':
            output = expect(client.get(job['download_url']), 200)
            assert_session(client, identity)
            print(f'PASS {mode}: {len(output.content)} bytes, {time.monotonic()-start:.1f}s, {response.http_version}', flush=True)
            return output.content, job
        time.sleep(1)
    raise AssertionError(f'{mode}: job exceeded {TIMEOUT:.0f}s')


def main():
    document = Document()
    document.add_heading('DocuVisa API verification', 0)
    document.add_paragraph('Lettre de motivation France Canada 2026')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Universite'
    table.cell(0, 1).text = 'Formation'
    word = io.BytesIO()
    document.save(word)

    presentation = Presentation()
    for title in ('DocuVisa API presentation', 'Deuxieme diapositive'):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
    pptx = io.BytesIO()
    presentation.save(pptx)

    with fitz.open() as pdf:
        for width, height in ((600, 800), (800, 400)):
            page = pdf.new_page(width=width, height=height)
            page.insert_text((50, 80), 'Mixed page verification', fontsize=18)
        mixed = pdf.tobytes()
    with fitz.open() as text_source:
        page = text_source.new_page(width=700, height=350)
        page.insert_text((50, 150), 'Bonjour Campus France 2026', fontsize=26)
        raster = page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes('png')
    with fitz.open() as pdf:
        page = pdf.new_page(width=700, height=350)
        page.insert_image(page.rect, stream=raster)
        scanned = pdf.tobytes()

    credentials = {'email': f'verification-documents-{uuid.uuid4().hex}@example.com', 'password': secrets.token_urlsafe(24), 'name': 'Verification documents'}
    with httpx.Client(base_url=ORIGIN, timeout=90, http2=HTTP2, headers={'Origin': ORIGIN}) as client:
        identity = expect(client.post('/api/auth/register', json=credentials), 200).json()['id']
        assert client.cookies.get('docuvisa_session'), 'Registration failed to set session cookie'
        assert_session(client, identity)

        pdf, _ = submit(client, identity, 'verification.docx', word.getvalue(), 'word-pdf')
        with fitz.open(stream=pdf, filetype='pdf') as output:
            assert len(output) == 1 and 'DocuVisa API verification' in output[0].get_text()
        docx, _ = submit(client, identity, 'verification.pdf', pdf, 'pdf-word')
        assert 'DocuVisa API verification' in ' '.join(p.text for p in Document(io.BytesIO(docx)).paragraphs)

        pdf, _ = submit(client, identity, 'verification.pptx', pptx.getvalue(), 'pptx-pdf')
        with fitz.open(stream=pdf, filetype='pdf') as output:
            assert len(output) == 2
            assert 'DocuVisa API presentation' in output[0].get_text()
            assert 'Deuxieme diapositive' in output[1].get_text()

        output, _ = submit(client, identity, 'mixed-pages.pdf', mixed, 'pdf-pptx')
        slides = Presentation(io.BytesIO(output)).slides
        assert len(slides) == 2
        for slide, ratio in zip(slides, (.75, 2)):
            picture = slide.shapes[0]
            assert abs(picture.width / picture.height - ratio) < .001
            with Image.open(io.BytesIO(picture.image.blob)) as image:
                assert image.width * image.height <= 4_010_000

        docx, job = submit(client, identity, 'scan.pdf', scanned, 'pdf-word')
        assert job['metadata']['ocr'] is True
        assert 'Campus France 2026' in ' '.join(p.text for p in Document(io.BytesIO(docx)).paragraphs)

        expect(client.post('/api/jobs', files={'file': ('invalid.pdf', b'not a PDF')}, data={'mode': 'pdf-word', 'options': '{}'}), 422)
        assert_session(client, identity)
        history = expect(client.get('/api/jobs'), 200).json()
        assert len(history) == 5 and all(job['status'] == 'completed' for job in history)
        expect(client.post('/api/auth/logout'), 200)
        assert expect(client.get('/api/auth/session'), 200).json()['user'] is None
        expect(client.post('/api/auth/login', json=credentials), 200)
        assert_session(client, identity)
        print('PASS invalid PDF422, archive history, logout/login and persistent sessions', flush=True)
    print('ALL DOCUMENT HTTP API CHECKS PASSED', flush=True)


if __name__ == '__main__':
    main()
